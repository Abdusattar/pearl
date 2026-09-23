"""Суммы по договорам Школы из Excel Айжан → скидки на карточках детей.

Тариф Школы один (18 000 с сентября 2026, решение владельца 23.09), у каждого
ребёнка сумма по договору другая (16 000 прошлогодним, 15 000 второй ребёнок,
9 000 учителям, 7 500 воспитателям, 6 000 / 12 000 по договору). Скидка =
тариф − сумма из таблицы, причина из колонки E. Идёт через children.set_discount —
тот же путь, что кнопка «Скидка» на карточке: аудит и пересчёт начисления
текущего месяца. Идемпотентно: повторный прогон ничего не меняет.

    python scripts/import_school_discounts.py "школа/Школа — суммы по детям 2026-27.xlsx" [--apply]

Без --apply — только показать, что изменится.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app.database import SessionLocal  # noqa: E402
from app.models import Student, User  # noqa: E402
from app.services import billing, children  # noqa: E402

SCHOOL_ORG_ID = 2
OWNER_USER_ID = 1


def main(path: str, apply: bool) -> None:
    ws = load_workbook(path, data_only=True)["Суммы по классам"]
    db = SessionLocal()
    owner = db.get(User, OWNER_USER_ID)
    changes, skipped = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cls, name, pin, total, reason = (list(row) + [None] * 5)[:5]
        if not pin:
            continue
        pin = str(pin).zfill(4)
        s = db.query(Student).filter(Student.organization_id == SCHOOL_ORG_ID, Student.pin == pin).first()
        if s is None:
            skipped.append((pin, name, "нет в системе"))
            continue
        try:
            total = float(str(total).replace(" ", "").replace(",", "."))
        except (TypeError, ValueError):
            skipped.append((pin, name, f"сумма не число: {total!r}"))
            continue
        base = billing.tuition_base_price(db, s)
        discount = round(base - total, 2)
        if discount < 0:
            skipped.append((pin, name, f"сумма {total:.0f} больше тарифа {base:.0f}"))
            continue
        reason = (str(reason).strip() if reason else "") or ("по договору" if discount > 0 else "")
        old = float(s.discount_amount or 0)
        if abs(old - discount) < 0.005 and (s.discount_reason or "") == (reason if discount > 0 else ""):
            continue
        changes.append((s, old, discount, reason))
    print(f"строк с изменением: {len(changes)}, пропущено: {len(skipped)}")
    for pin, name, why in skipped:
        print(f"  ПРОПУСК {pin} {name}: {why}")
    for s, old, new, reason in changes[:40]:
        print(f"  {s.pin} {s.name}: скидка {old:.0f} → {new:.0f} ({reason or 'без причины'})")
    if len(changes) > 40:
        print(f"  … и ещё {len(changes) - 40}")
    if not apply:
        print("Без --apply ничего не записано.")
        return
    recomputed = 0
    for s, _old, new, reason in changes:
        if children.set_discount(db, user=owner, student=s, amount=new, reason=reason):
            recomputed += 1
    db.commit()
    print(f"Записано {len(changes)} скидок, начислений текущего месяца пересчитано: {recomputed}.")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    main(args[0] if args else "школа/Школа — суммы по детям 2026-27.xlsx", "--apply" in sys.argv)
